# backend/apps/innovation/exportadores.py

import html
import io
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COLOR_OSCURO = "1A1A2E"
COLOR_ACENTO = "0F3460"
COLOR_ZEBRA = "F4F5F8"
MAX_COLUMNAS_VERTICAL = 5
KPIS_POR_FILA = 4

TIPOS_MIME = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "html": "text/html; charset=utf-8",
}
ALIAS_FORMATO = {"excel": "xlsx", "xls": "xlsx", "xlsx": "xlsx", "pdf": "pdf", "html": "html"}


class FormatoInvalido(Exception):
    """El formato pedido no es pdf, xlsx (excel) ni html."""


# ---------------------------------------------------------------- utilidades
def formatear(valor, tipo):
    """Texto de un valor según su tipo: dinero 'Bs 1,234.50', número, texto."""
    if valor is None:
        return ""
    if tipo == "dinero":
        return f"Bs {float(valor):,.2f}"
    if tipo == "numero":
        numero = float(valor)
        return str(int(numero)) if numero.is_integer() else f"{numero:.2f}"
    return str(valor)


def _nombre_archivo(titulo, generado, extension):
    """Nombre de archivo sin tildes ni espacios, con la fecha de generación."""
    limpio = unicodedata.normalize("NFD", str(titulo or "reporte"))
    limpio = "".join(c for c in limpio if unicodedata.category(c) != "Mn").lower()
    base = re.sub(r"[^a-z0-9]+", "_", limpio).strip("_")[:50] or "reporte"
    m = re.match(r"(\d\d)/(\d\d)/(\d{4}) (\d\d):(\d\d)", str(generado or ""))
    sello = f"{m.group(3)}{m.group(2)}{m.group(1)}_{m.group(4)}{m.group(5)}" if m else "reporte"
    return f"{base}_{sello}.{extension}"


def _lineas_cabecera(resultado, meta):
    """Textos de contexto que van bajo el título: periodo, pregunta y fecha."""
    lineas = []
    if resultado.get("subtitulo"):
        lineas.append(resultado["subtitulo"])
    if meta.get("pregunta"):
        lineas.append(f"Pregunta: {meta['pregunta']}")
    if meta.get("generado"):
        lineas.append(f"Generado: {meta['generado']}")
    return lineas


# ---------------------------------------------------------------- HTML
def exportar_html(resultado, meta):
    """Documento HTML autocontenido (estilos incluidos, sin recursos externos)."""
    e = html.escape
    columnas = resultado["columnas"]
    partes = [
        "<!DOCTYPE html>",
        '<html lang="es">',
        "<head>",
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        f"<title>{e(resultado['titulo'])} - SmartLook</title>",
        "<style>",
        "body{font-family:Arial,Helvetica,sans-serif;margin:32px;color:#1e1a16;background:#fff}",
        f"h1{{margin:0 0 4px;color:#{COLOR_OSCURO}}}",
        ".marca{color:#6b6459;font-size:13px;margin-bottom:16px}",
        ".ctx{color:#555;font-size:14px;margin:2px 0}",
        ".kpis{display:flex;flex-wrap:wrap;gap:12px;margin:20px 0}",
        ".kpi{border:1px solid #dcdde4;border-radius:10px;padding:10px 18px;min-width:150px}",
        ".kpi span{display:block;font-size:12px;color:#6b6459}",
        f".kpi strong{{display:block;font-size:20px;color:#{COLOR_ACENTO};margin-top:2px}}",
        "table{border-collapse:collapse;width:100%;margin-top:8px;font-size:14px}",
        f"th{{background:#{COLOR_OSCURO};color:#fff;text-align:left;padding:9px 10px}}",
        "td{padding:8px 10px;border-bottom:1px solid #e5e6ec}",
        f"tbody tr:nth-child(even) td{{background:#{COLOR_ZEBRA}}}",
        ".num{text-align:right}",
        ".nota{margin-top:14px;padding:10px 14px;background:#fff8e6;border-left:4px solid #ff9800;font-size:14px}",
        ".vacio{margin-top:16px;color:#6b6459}",
        ".pie{margin-top:28px;font-size:12px;color:#8a8fa0}",
        "</style>",
        "</head>",
        "<body>",
        f"<h1>{e(resultado['titulo'])}</h1>",
        '<div class="marca">SmartLook · Reporte de gestión</div>',
    ]
    for linea in _lineas_cabecera(resultado, meta):
        partes.append(f'<p class="ctx">{e(linea)}</p>')

    if resultado["kpis"]:
        partes.append('<div class="kpis">')
        for k in resultado["kpis"]:
            partes.append(
                f'<div class="kpi"><span>{e(k["etiqueta"])}</span>'
                f'<strong>{e(formatear(k["valor"], k["tipo"]))}</strong></div>'
            )
        partes.append("</div>")

    if resultado["filas"]:
        partes.append("<table><thead><tr>")
        for c in columnas:
            clase = ' class="num"' if c["tipo"] in ("numero", "dinero") else ""
            partes.append(f"<th{clase}>{e(c['nombre'])}</th>")
        partes.append("</tr></thead><tbody>")
        for fila in resultado["filas"]:
            partes.append("<tr>")
            for c, valor in zip(columnas, fila):
                clase = ' class="num"' if c["tipo"] in ("numero", "dinero") else ""
                partes.append(f"<td{clase}>{e(formatear(valor, c['tipo']))}</td>")
            partes.append("</tr>")
        partes.append("</tbody></table>")
    else:
        partes.append('<p class="vacio">Sin filas para mostrar.</p>')

    for nota in resultado["notas"]:
        partes.append(f'<div class="nota">{e(nota)}</div>')
    partes.append('<p class="pie">Generado con el asistente de gestión de SmartLook.</p>')
    partes.extend(["</body>", "</html>"])
    return "\n".join(partes).encode("utf-8")


# ---------------------------------------------------------------- Excel
def exportar_excel(resultado, meta):
    """Libro de Excel con cabecera, indicadores y la tabla con formato numérico."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    negrita = Font(bold=True)
    relleno = PatternFill("solid", fgColor=COLOR_OSCURO)
    borde = Border(bottom=Side(style="thin", color="DDDDDD"))

    ws.append([resultado["titulo"]])
    ws["A1"].font = Font(bold=True, size=15, color=COLOR_OSCURO)
    for linea in _lineas_cabecera(resultado, meta):
        ws.append([linea])
    ws.append([])

    if resultado["kpis"]:
        ws.append(["Indicador", "Valor"])
        fila_cab = ws.max_row
        for celda in ws[fila_cab][:2]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno
        for k in resultado["kpis"]:
            ws.append([k["etiqueta"], k["valor"]])
            celda = ws.cell(row=ws.max_row, column=2)
            celda.number_format = _formato_excel(k["tipo"], k["valor"])
            celda.alignment = Alignment(horizontal="right")
        ws.append([])

    columnas = resultado["columnas"]
    ws.append([c["nombre"] for c in columnas])
    fila_tabla = ws.max_row
    for celda in ws[fila_tabla]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila in resultado["filas"]:
        ws.append(list(fila))
        for indice, c in enumerate(columnas, start=1):
            celda = ws.cell(row=ws.max_row, column=indice)
            celda.border = borde
            if c["tipo"] in ("numero", "dinero"):
                celda.number_format = _formato_excel(c["tipo"], celda.value)
                celda.alignment = Alignment(horizontal="right")
    if resultado["filas"]:
        ws.freeze_panes = ws.cell(row=fila_tabla + 1, column=1)
        ws.auto_filter.ref = (
            f"A{fila_tabla}:{get_column_letter(len(columnas))}{ws.max_row}"
        )
    else:
        ws.append(["Sin filas para mostrar."])

    for nota in resultado["notas"]:
        ws.append([])
        ws.append([f"Nota: {nota}"])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="B26A00")

    # Ancho de columnas según el contenido de la tabla (sin contar el título).
    for indice in range(1, max(len(columnas), 2) + 1):
        letra = get_column_letter(indice)
        ancho = 10
        for fila in ws.iter_rows(min_row=fila_tabla, min_col=indice, max_col=indice):
            for celda in fila:
                if celda.value is not None:
                    ancho = max(ancho, min(len(str(celda.value)) + 3, 45))
        ws.column_dimensions[letra].width = ancho
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 10, 28)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _formato_excel(tipo, valor):
    """Formato numérico de Excel según el tipo del dato."""
    if tipo == "dinero":
        return '"Bs" #,##0.00'
    if isinstance(valor, float) and not float(valor).is_integer():
        return "0.00"
    return "0"


# ---------------------------------------------------------------- PDF
def exportar_pdf(resultado, meta):
    """PDF A4 (horizontal si la tabla es ancha) con indicadores y tabla paginada."""
    columnas = resultado["columnas"]
    horizontal = len(columnas) > MAX_COLUMNAS_VERTICAL
    pagina = landscape(A4) if horizontal else A4
    ancho_util = pagina[0] - 3 * cm

    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("Titulo", parent=estilos["Title"], alignment=0,
                            textColor=colors.HexColor(f"#{COLOR_OSCURO}"), fontSize=20, spaceAfter=2)
    marca = ParagraphStyle("Marca", parent=estilos["Normal"], fontSize=9,
                           textColor=colors.HexColor("#6B6459"), spaceAfter=8)
    contexto = ParagraphStyle("Contexto", parent=estilos["Normal"], fontSize=10,
                              textColor=colors.HexColor("#444444"), leading=13)
    celda = ParagraphStyle("Celda", parent=estilos["Normal"], fontSize=8.5, leading=10.5)
    celda_num = ParagraphStyle("CeldaNum", parent=celda, alignment=2)
    cabecera = ParagraphStyle("Cabecera", parent=celda, textColor=colors.white,
                              fontName="Helvetica-Bold")
    cabecera_num = ParagraphStyle("CabeceraNum", parent=cabecera, alignment=2)
    nota_estilo = ParagraphStyle("Nota", parent=estilos["Normal"], fontSize=9,
                                 textColor=colors.HexColor("#8A5A00"), leading=12, spaceBefore=6)
    esc = html.escape

    elementos = [
        Paragraph(esc(resultado["titulo"]), titulo),
        Paragraph("SmartLook · Reporte de gestión", marca),
    ]
    for linea in _lineas_cabecera(resultado, meta):
        elementos.append(Paragraph(esc(linea), contexto))
    elementos.append(Spacer(1, 0.4 * cm))

    if resultado["kpis"]:
        kpi_estilo_e = ParagraphStyle("KpiE", parent=celda, textColor=colors.HexColor("#6B6459"), fontSize=8)
        kpi_estilo_v = ParagraphStyle("KpiV", parent=celda, fontSize=13, leading=16,
                                      fontName="Helvetica-Bold",
                                      textColor=colors.HexColor(f"#{COLOR_ACENTO}"))
        bloques = []
        for k in resultado["kpis"]:
            bloques.append([
                Paragraph(esc(k["etiqueta"]), kpi_estilo_e),
                Paragraph(esc(formatear(k["valor"], k["tipo"])), kpi_estilo_v),
            ])
        filas_kpi = [bloques[i:i + KPIS_POR_FILA] for i in range(0, len(bloques), KPIS_POR_FILA)]
        for fila in filas_kpi:
            while len(fila) < KPIS_POR_FILA:
                fila.append("")
        tabla_kpi = Table(filas_kpi, colWidths=[ancho_util / KPIS_POR_FILA] * KPIS_POR_FILA)
        tabla_kpi.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DCDDE4")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DCDDE4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        elementos.extend([tabla_kpi, Spacer(1, 0.5 * cm)])

    if resultado["filas"]:
        numericas = [c["tipo"] in ("numero", "dinero") for c in columnas]
        datos = [[
            Paragraph(esc(c["nombre"]), cabecera_num if num else cabecera)
            for c, num in zip(columnas, numericas)
        ]]
        pesos = [max(6, min(len(c["nombre"]), 30)) for c in columnas]
        for fila in resultado["filas"]:
            renglon = []
            for indice, (c, valor) in enumerate(zip(columnas, fila)):
                texto = formatear(valor, c["tipo"])
                pesos[indice] = max(pesos[indice], min(len(texto), 34))
                renglon.append(Paragraph(esc(texto), celda_num if numericas[indice] else celda))
            datos.append(renglon)
        total = float(sum(pesos))
        anchos = [ancho_util * p / total for p in pesos]
        tabla = Table(datos, colWidths=anchos, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{COLOR_OSCURO}")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{COLOR_ZEBRA}")]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E6EC")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elementos.append(tabla)
    else:
        elementos.append(Paragraph("Sin filas para mostrar.", contexto))

    for nota in resultado["notas"]:
        elementos.append(Paragraph(f"Nota: {esc(nota)}", nota_estilo))

    generado = meta.get("generado", "")

    def pie(lienzo, documento):
        lienzo.saveState()
        lienzo.setFont("Helvetica", 8)
        lienzo.setFillColor(colors.HexColor("#8A8FA0"))
        lienzo.drawString(1.5 * cm, 0.9 * cm, f"SmartLook · {generado}")
        lienzo.drawRightString(pagina[0] - 1.5 * cm, 0.9 * cm, f"Página {documento.page}")
        lienzo.restoreState()

    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer, pagesize=pagina, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.6 * cm, title=resultado["titulo"], author="SmartLook",
    )
    documento.build(elementos, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


# ---------------------------------------------------------------- entrada
def exportar(formato, resultado, meta):
    """Genera el archivo pedido. Devuelve (contenido, tipo_mime, nombre_archivo)."""
    clave = ALIAS_FORMATO.get(str(formato or "").strip().lower())
    if clave is None:
        raise FormatoInvalido(f"Formato no válido: {formato}. Usa pdf, xlsx o html.")
    meta = meta or {}
    generadores = {"pdf": exportar_pdf, "xlsx": exportar_excel, "html": exportar_html}
    contenido = generadores[clave](resultado, meta)
    nombre = _nombre_archivo(resultado.get("titulo"), meta.get("generado"), clave)
    return contenido, TIPOS_MIME[clave], nombre